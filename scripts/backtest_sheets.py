"""Backtest and Edge Sweep sheets (discovery seasons only). Called from build_workbook.py."""
import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

DATA_COLS = (["season", "window", "date", "home", "away"]
             + [f"h_{k}" for k in ("gp", "cum_pts", "cum_opp", "cum_poss", "avg_pace", "prior_o", "prior_d", "prior_pace", "rest", "g3in4")]
             + [f"a_{k}" for k in ("gp", "cum_pts", "cum_opp", "cum_poss", "avg_pace", "prior_o", "prior_d", "prior_pace", "rest", "g3in4")]
             + ["lg_ortg", "book", "open_home_spread", "mkt_home_spread", "open_total", "mkt_total",
               "ml_home_open", "ml_home_close", "ml_away_open", "ml_away_close", "home_pts", "away_pts"])
FORM_COLS = ["h_bo", "h_bd", "h_bp", "a_bo", "a_bd", "a_bp", "pace", "hpts", "apts", "restadj", "fatadj", "margin", "total",
             "act_margin", "act_total", "s_edge", "s_side", "s_res", "s_abs", "t_edge", "t_side", "t_res", "t_abs",
             "mod_err", "mod_abs", "mkt_abs", "tot_mod_abs", "tot_mkt_abs",
             "o_edge", "o_side", "o_res", "o_abs", "o_clv", "to_edge", "to_side", "to_res", "to_abs", "to_clv"]
NAMES = DATA_COLS + FORM_COLS
L = {n: get_column_letter(i + 1) for i, n in enumerate(NAMES)}

LABELS = {
    "season": "Season", "window": "Window", "date": "Date", "home": "Home", "away": "Away",
    "lg_ortg": "League ORtg (pre-game)", "book": "Reference book", "open_home_spread": "Home spread OPEN", "mkt_home_spread": "Home spread CLOSE",
    "open_total": "Total OPEN", "mkt_total": "Total CLOSE", "ml_home_open": "ML home open", "ml_home_close": "ML home close",
    "ml_away_open": "ML away open", "ml_away_close": "ML away close", "home_pts": "Home PTS", "away_pts": "Away PTS",
    "h_bo": "Home blended ORtg", "h_bd": "Home blended DRtg", "h_bp": "Home blended pace",
    "a_bo": "Away blended ORtg", "a_bd": "Away blended DRtg", "a_bp": "Away blended pace", "pace": "Expected pace",
    "hpts": "Home proj pts", "apts": "Away proj pts", "restadj": "Rest adj", "fatadj": "3-in-4 adj",
    "margin": "Proj home margin", "total": "Proj total", "act_margin": "Actual margin", "act_total": "Actual total",
    "s_edge": "Spread edge", "s_side": "Spread side", "s_res": "Spread result", "s_abs": "|Spread edge|",
    "t_edge": "Total edge", "t_side": "Total side", "t_res": "Total result", "t_abs": "|Total edge|",
    "mod_err": "Model error (proj - actual)", "mod_abs": "Model |margin err|", "mkt_abs": "Market |margin err|",
    "tot_mod_abs": "Model |total err|", "tot_mkt_abs": "Market |total err|",
    "o_edge": "Spread edge vs OPEN", "o_side": "Side (open)", "o_res": "Result at open line", "o_abs": "|Edge vs open|", "o_clv": "CLV pts (bet at open)",
    "to_edge": "Total edge vs OPEN", "to_side": "Total side (open)", "to_res": "Total result at open", "to_abs": "|Total edge vs open|", "to_clv": "Total CLV pts (bet at open)"}
_SUB = {"gp": "GP (pre)", "cum_pts": "cum PTS", "cum_opp": "cum Opp PTS", "cum_poss": "cum poss", "avg_pace": "avg pace",
        "prior_o": "prior-season ORtg", "prior_d": "prior-season DRtg", "prior_pace": "prior-season pace",
        "rest": "rest days (cap 3)", "g3in4": "3rd game in 4 nights"}
for n in NAMES:
    if n not in LABELS:
        side, k = n.split("_", 1)
        LABELS[n] = ("Home " if side == "h" else "Away ") + _SUB[k]


def add(wb, pit, h):
    """h: dict of helpers from build_workbook (header, section, fonts)."""
    header, section, title, note, blue = h["header"], h["section"], h["title"], h["note"], h["blue"]
    wbk = wb.create_sheet("Backtest", 5)
    header(wbk, 1, [LABELS[n] for n in NAMES]); wbk.row_dimensions[1].height = 54
    K, HCA, REST, CAP, PEN = ("Assumptions!$B$9", "Assumptions!$B$10", "Assumptions!$B$16",
                              "Assumptions!$B$17", "Assumptions!$B$18")
    for i, rec in enumerate(pit.to_dict("records"), 2):
        for n in DATA_COLS:
            v = rec[n]
            wbk.cell(row=i, column=NAMES.index(n) + 1, value=None if pd.isna(v) else v)
        c = lambda n: f"{L[n]}{i}"
        F = {}
        for p in ("h", "a"):
            gp = c(p + "_gp")
            F[p + "_bo"] = f"=({K}*{c(p+'_prior_o')}+{gp}*IF({gp}=0,0,100*{c(p+'_cum_pts')}/{c(p+'_cum_poss')}))/({K}+{gp})"
            F[p + "_bd"] = f"=({K}*{c(p+'_prior_d')}+{gp}*IF({gp}=0,0,100*{c(p+'_cum_opp')}/{c(p+'_cum_poss')}))/({K}+{gp})"
            F[p + "_bp"] = f"=({K}*{c(p+'_prior_pace')}+{gp}*{c(p+'_avg_pace')})/({K}+{gp})"
        F["pace"] = f"=({c('h_bp')}+{c('a_bp')})/2"
        F["hpts"] = f"={c('pace')}*({c('h_bo')}+{c('a_bd')}-{c('lg_ortg')})/100"
        F["apts"] = f"={c('pace')}*({c('a_bo')}+{c('h_bd')}-{c('lg_ortg')})/100"
        F["restadj"] = f"={REST}*MAX(-{CAP},MIN({CAP},{c('h_rest')}-{c('a_rest')}))"
        F["fatadj"] = f"={c('h_g3in4')}*{PEN}-{c('a_g3in4')}*{PEN}"
        F["margin"] = f"={c('hpts')}-{c('apts')}+{HCA}+{c('restadj')}+{c('fatadj')}"
        F["total"] = f"={c('hpts')}+{c('apts')}"
        F["act_margin"] = f"={c('home_pts')}-{c('away_pts')}"
        F["act_total"] = f"={c('home_pts')}+{c('away_pts')}"
        F["s_edge"] = f"={c('margin')}+{c('mkt_home_spread')}"
        F["s_side"] = f'=IF({c("s_edge")}>0,"Home","Away")'
        cover = f'{c("act_margin")}+{c("mkt_home_spread")}'
        F["s_res"] = f'=IF({cover}=0,"Push",IF(({c("s_side")}="Home")=({cover}>0),"Win","Loss"))'
        F["s_abs"] = f"=ABS({c('s_edge')})"
        F["t_edge"] = f"={c('total')}-{c('mkt_total')}"
        F["t_side"] = f'=IF({c("t_edge")}>0,"Over","Under")'
        F["t_res"] = f'=IF({c("act_total")}={c("mkt_total")},"Push",IF(({c("t_side")}="Over")=({c("act_total")}>{c("mkt_total")}),"Win","Loss"))'
        F["t_abs"] = f"=ABS({c('t_edge')})"
        F["mod_err"] = f"={c('margin')}-{c('act_margin')}"
        F["mod_abs"] = f"=ABS({c('mod_err')})"
        F["mkt_abs"] = f"=ABS(-{c('mkt_home_spread')}-{c('act_margin')})"
        F["tot_mod_abs"] = f"=ABS({c('total')}-{c('act_total')})"
        F["tot_mkt_abs"] = f"=ABS({c('mkt_total')}-{c('act_total')})"
        F["o_edge"] = f"={c('margin')}+{c('open_home_spread')}"
        F["o_side"] = f'=IF({c("o_edge")}>0,"Home","Away")'
        cov_o = f'{c("act_margin")}+{c("open_home_spread")}'
        F["o_res"] = f'=IF({cov_o}=0,"Push",IF(({c("o_side")}="Home")=({cov_o}>0),"Win","Loss"))'
        F["o_abs"] = f"=ABS({c('o_edge')})"
        F["o_clv"] = f'=IF({c("o_side")}="Home",1,-1)*({c("open_home_spread")}-{c("mkt_home_spread")})'
        F["to_edge"] = f"={c('total')}-{c('open_total')}"
        F["to_side"] = f'=IF({c("to_edge")}>0,"Over","Under")'
        F["to_res"] = f'=IF({c("act_total")}={c("open_total")},"Push",IF(({c("to_side")}="Over")=({c("act_total")}>{c("open_total")}),"Win","Loss"))'
        F["to_abs"] = f"=ABS({c('to_edge')})"
        F["to_clv"] = f'=IF({c("to_side")}="Over",1,-1)*({c("mkt_total")}-{c("open_total")})'
        for n in FORM_COLS:
            wbk.cell(row=i, column=NAMES.index(n) + 1, value=F[n])
    NB = len(pit) + 1
    wbk.freeze_panes = "F2"
    for j in range(1, len(NAMES) + 1): wbk.column_dimensions[get_column_letter(j)].width = 11

    ws = wb.create_sheet("Edge Sweep", 6)
    ws["A1"] = "Edge Sweep (Backtest, discovery seasons only)"; ws["A1"].font = title
    ws["A2"] = ("Every game is bet on the side the model prefers; the threshold keeps only games with |edge| >= threshold. Lines are ESPN "
                "open/close from one book (ESPN BET; DraftKings for 4 games). 'Graded at CLOSE' tables test win rate against the closing line; "
                "'bet at OPEN' tables grade at the opening line and report average closing-line value (CLV). A real signal should improve as "
                "the threshold rises; a flat curve means the gap is noise. Placeholder parameters; discovery seasons only.")
    ws["A2"].font = note; ws["A2"].alignment = Alignment(wrap_text=True); ws.row_dimensions[2].height = 40
    ws["A4"] = "Season filter (All, 2023-24 or 2024-25)"; ws["B4"] = "All"; ws["B4"].font = blue
    ws["A5"] = "Season criterion (helper)"; ws["B5"] = '=IF(B4="All","*",B4)'
    ws["A6"] = "Window (Discovery only; Confirmation rows are not built)"; ws["B6"] = "Discovery"; ws["B6"].font = blue
    ws["A7"] = "Breakeven win rate at standard price"
    ws["B7"] = "=IF(Assumptions!B24<0,-Assumptions!B24/(-Assumptions!B24+100),100/(Assumptions!B24+100))"; ws["B7"].number_format = "0.00%"
    ws["A8"] = "Win payout per 1 unit staked"
    ws["B8"] = "=IF(Assumptions!B24<0,100/-Assumptions!B24,Assumptions!B24/100)"; ws["B8"].number_format = "0.0000"
    R = lambda n: f"Backtest!${L[n]}$2:${L[n]}${NB}"
    crit = f"{R('window')},$B$6,{R('season')},$B$5"
    section(ws, 10, "ACCURACY (filtered games unless noted)", 4)
    acc = [
        (11, "Games", f"=COUNTIFS({crit})", "0"),
        (12, "Model |margin error| (pts)", f"=AVERAGEIFS({R('mod_abs')},{crit})", "0.00"),
        (13, "Market |margin error| (pts)", f"=AVERAGEIFS({R('mkt_abs')},{crit})", "0.00"),
        (14, "Model |total error| (pts)", f"=AVERAGEIFS({R('tot_mod_abs')},{crit})", "0.00"),
        (15, "Market |total error| (pts)", f"=AVERAGEIFS({R('tot_mkt_abs')},{crit})", "0.00"),
        (16, "Model bias (proj - actual margin)", f"=AVERAGEIFS({R('mod_err')},{crit})", "0.00"),
        (17, "Correlation: model margin vs actual (all rows)", f"=CORREL({R('margin')},{R('act_margin')})", "0.000"),
        (18, "Correlation: market margin vs actual (all rows)", f"=CORREL({R('mkt_home_spread')},{R('act_margin')})*-1", "0.000"),
    ]
    for r, a, f, nf in acc:
        ws.cell(row=r, column=1, value=a); ws.cell(row=r, column=2, value=f).number_format = nf
    thr = [0, 0.5, 1, 1.5, 2, 2.5, 3, 4, 5, 6, 8]

    def sweep(top, label, abs_n, res_n, clv_n=None):
        section(ws, top, label, 11 if clv_n else 10)
        header(ws, top + 1, ["Min |edge| (pts)", "Picks", "Wins", "Losses", "Pushes", "Win rate", "z vs breakeven",
                             "p (one-sided)", "Units won", "ROI per bet"] + (["Avg CLV (pts) vs close"] if clv_n else []))
        for k, t in enumerate(thr):
            r = top + 2 + k
            ws.cell(row=r, column=1, value=t).font = blue
            cr = f'{crit},{R(abs_n)},">="&$A{r}'
            ws.cell(row=r, column=2, value=f"=COUNTIFS({cr})")
            for col, res in ((3, "Win"), (4, "Loss"), (5, "Push")):
                ws.cell(row=r, column=col, value=f'=COUNTIFS({cr},{R(res_n)},"{res}")')
            ws.cell(row=r, column=6, value=f"=IF(C{r}+D{r}=0,0,C{r}/(C{r}+D{r}))").number_format = "0.0%"
            ws.cell(row=r, column=7, value=f"=IF(C{r}+D{r}=0,0,(F{r}-$B$7)/SQRT($B$7*(1-$B$7)/(C{r}+D{r})))").number_format = "0.00"
            ws.cell(row=r, column=8, value=f"=1-NORMSDIST(G{r})").number_format = "0.000"
            ws.cell(row=r, column=9, value=f"=C{r}*$B$8-D{r}").number_format = "0.0"
            ws.cell(row=r, column=10, value=f"=IF(C{r}+D{r}+E{r}=0,0,I{r}/(C{r}+D{r}+E{r}))").number_format = "0.0%"
            if clv_n:
                ws.cell(row=r, column=11, value=f"=IF(B{r}=0,0,AVERAGEIFS({R(clv_n)},{cr}))").number_format = "0.00"

    sweep(21, "SPREAD sweep (graded at CLOSE)", "s_abs", "s_res")
    sweep(36, "TOTAL sweep (graded at CLOSE)", "t_abs", "t_res")
    sweep(51, "SPREAD, bet at OPEN (graded at open line; CLV = movement toward the pick)", "o_abs", "o_res", "o_clv")
    sweep(66, "TOTAL, bet at OPEN (graded at open line; CLV = movement toward the pick)", "to_abs", "to_res", "to_clv")
    ws.column_dimensions["A"].width = 52
    for j in range(2, 12): ws.column_dimensions[get_column_letter(j)].width = 13
